"""Document classifier (BE-05): rules first, LLM fallback.

Contract §5 (doc_type), §6 (documents). The rules need only the filename, the
MIME type and the text of the first pages, so most uploads never touch the LLM.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import PurePosixPath

from pydantic import BaseModel, Field

from app.llm.client import complete_json
from app.logging import get_logger
from app.models.enums import DocType

logger = get_logger(__name__)

PEEK_MAX_CHARS = 6000
LLM_MIN_TEXT_CHARS = 20
LLM_MAX_TEXT_CHARS = 3000

_TABULAR_EXTENSIONS = {".csv", ".tsv", ".xlsx", ".xls"}
_TABULAR_MIMES = {
    "text/csv",
    "application/vnd.ms-excel",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

# Header names for the survey rule (same variants BE-06's tabular parser accepts).
_MD_NAMES = {"md", "measured depth", "depth md"}
_INC_NAMES = {"inc", "incl", "inclination"}
_AZI_NAMES = {"azi", "azim", "azimuth"}

_XML_ROOT_RE = re.compile(r"<(?![?!])(?:[\w.-]+:)?([\w.-]+)")
_WITSML_ROOTS = {"drillreports", "drillreport"}
_LAS_START_RE = re.compile(r"\s*~V", re.IGNORECASE)
_LAS_DATA_RE = re.compile(r"^~A", re.IGNORECASE | re.MULTILINE)

# Order matters: the first matching rule wins (PRD BE-05).
_KEYWORD_RULES: tuple[tuple[DocType, re.Pattern[str]], ...] = (
    (DocType.DDR, re.compile(r"daily\s+drilling\s+report|\biadc\b|morning\s+report", re.IGNORECASE)),
    (DocType.WCR, re.compile(r"well\s+completion\s+report|completion\s+report", re.IGNORECASE)),
    (DocType.MUD_LOG, re.compile(r"mud\s*log|master\s*log", re.IGNORECASE)),
    (DocType.CEMENT_REPORT, re.compile(r"cement\s+job|cementing\s+report", re.IGNORECASE)),
    (
        DocType.PROGRAM,
        re.compile(r"drilling\s+program(?:me)?|casing\s+program(?:me)?|mud\s+program(?:me)?", re.IGNORECASE),
    ),
    (DocType.INCIDENT, re.compile(r"\bincident\b|\bnpt\s+report\b|\bfishing\b", re.IGNORECASE)),
)

_WELL_NAME_PATTERNS = (
    re.compile(r"<nameWell[^>]*>\s*([^<]+?)\s*</nameWell>", re.IGNORECASE),
    re.compile(r"^[ \t]*well(?:[ \t]*name)?[ \t]*[:=-][ \t]*([A-Za-z0-9][\w\-./ ]{0,59}?)[ \t]*$", re.IGNORECASE | re.MULTILINE),
)

_LLM_SYSTEM = (
    "You classify oil-well documents from the filename and the text of their first pages.\n"
    "doc_type must be one of: wcr (well completion report), ddr (daily drilling report), "
    "mud_log, program (drilling/casing/mud program), cement_report, incident (incident or NPT "
    "report), survey (deviation survey table), las, witsml, other.\n"
    "well_name is the well's name exactly as written in the header, or null if it is not stated. "
    "Never guess. confidence is your 0-1 estimate that doc_type is right."
)


class _LlmClassification(BaseModel):
    doc_type: DocType
    well_name: str | None = None
    confidence: float = Field(ge=0, le=1)


@dataclass(frozen=True)
class Classification:
    doc_type: DocType
    well_name: str | None = None
    confidence: float | None = None
    method: str = "rules"  # 'requested' | 'rules' | 'llm' | 'none'


def peek_text(data: bytes, filename: str, mime: str | None, max_chars: int = PEEK_MAX_CHARS) -> str:
    """Text of the first (up to) two pages, cheaply and without OCR.

    PDFs use the text layer when PyMuPDF is installed; scans, images and
    unreadable files return "" (classification then falls back to 'other').
    """
    ext = PurePosixPath(filename.lower()).suffix
    if ext == ".pdf" or mime == "application/pdf":
        return _peek_pdf(data, max_chars)
    if ext in {".png", ".jpg", ".jpeg", ".tif", ".tiff"} or (mime or "").startswith("image/"):
        return ""
    if ext in {".xlsx", ".xls"}:
        return _peek_xlsx(data, max_chars)
    sample = data[:max_chars * 4]
    if b"\x00" in sample:
        return ""
    return sample.decode("utf-8", errors="replace")[:max_chars]


def _peek_pdf(data: bytes, max_chars: int) -> str:
    try:
        import fitz  # PyMuPDF; installed with BE-07

        with fitz.open(stream=data, filetype="pdf") as pdf:
            text = "\n".join(pdf[i].get_text() for i in range(min(2, len(pdf))))
        return text[:max_chars]
    except Exception:  # noqa: BLE001 - missing dependency or a broken PDF both mean "no text"
        logger.info("pdf_peek_unavailable")
        return ""


def _peek_xlsx(data: bytes, max_chars: int) -> str:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        lines = []
        for row in sheet.iter_rows(min_row=1, max_row=20, values_only=True):
            lines.append(",".join("" if cell is None else str(cell) for cell in row))
        return "\n".join(lines)[:max_chars]
    except Exception:  # noqa: BLE001
        logger.info("xlsx_peek_unavailable")
        return ""


def guess_well_name(text: str) -> str | None:
    """Well name from a WITSML <nameWell> element or a 'Well: …' header line."""
    for pattern in _WELL_NAME_PATTERNS:
        match = pattern.search(text)
        if match:
            name = match.group(1).strip()
            if name:
                return name
    return None


def _is_witsml(text: str) -> bool:
    stripped = text.lstrip()
    if not stripped.startswith("<"):
        return False
    if "witsml.org" in stripped[:2000].lower():
        return True
    match = _XML_ROOT_RE.search(stripped)
    return bool(match) and match.group(1).lower() in _WITSML_ROOTS


def _is_las(text: str) -> bool:
    return bool(_LAS_START_RE.match(text)) or bool(_LAS_DATA_RE.search(text))


def _normalise_header(cell: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"\(.*?\)", "", cell)).strip().lower()


def _is_survey_table(filename: str, mime: str | None, text: str) -> bool:
    ext = PurePosixPath(filename.lower()).suffix
    if ext not in _TABULAR_EXTENSIONS and mime not in _TABULAR_MIMES:
        return False
    first_line = next((line for line in text.splitlines() if line.strip()), "")
    headers = {_normalise_header(cell) for cell in re.split(r"[,;\t|]", first_line)}
    return bool(headers & _MD_NAMES) and bool(headers & _INC_NAMES) and bool(headers & _AZI_NAMES)


def classify_by_rules(filename: str, first_page_text: str, mime: str | None) -> DocType | None:
    """Deterministic rules from PRD BE-05; None when nothing matches."""
    if _is_witsml(first_page_text):
        return DocType.WITSML
    if _is_las(first_page_text):
        return DocType.LAS
    if _is_survey_table(filename, mime, first_page_text):
        return DocType.SURVEY
    for doc_type, pattern in _KEYWORD_RULES:
        if pattern.search(first_page_text):
            return doc_type
    return None


async def classify_with_details(
    filename: str,
    first_page_text: str,
    mime: str | None,
    requested: DocType | None = None,
) -> Classification:
    if requested is not None:
        return Classification(requested, method="requested")

    by_rules = classify_by_rules(filename, first_page_text, mime)
    if by_rules is not None:
        return Classification(by_rules, method="rules")

    text = first_page_text.strip()
    if len(text) < LLM_MIN_TEXT_CHARS:
        # nothing for the LLM to read (e.g. a scan before OCR); don't ask it to guess
        return Classification(DocType.OTHER, method="none")

    user = f"Filename: {filename}\n\nFirst pages:\n{text[:LLM_MAX_TEXT_CHARS]}"
    try:
        result, _meta = await complete_json(_LLM_SYSTEM, user, _LlmClassification, max_tokens=200)
    except Exception:  # noqa: BLE001 - a failed classification must not fail the upload
        logger.exception("classify_llm_failed filename=%s", filename)
        return Classification(DocType.OTHER, method="none")
    return Classification(result.doc_type, result.well_name, result.confidence, method="llm")


async def classify(
    filename: str,
    first_page_text: str,
    mime: str | None,
    requested: DocType | None = None,
) -> DocType:
    """`classify(filename, first_page_text, mime) -> doc_type` (PRD BE-05)."""
    return (await classify_with_details(filename, first_page_text, mime, requested)).doc_type
